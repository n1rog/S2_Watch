"""Write a MAGNET S2 Intelligence Report as a Word .docx in the MAGNET format.

Arial throughout; US Letter, 0.75" margins; tables 10080 DXA wide; a navy title block;
the navy/off-white routing table; blue all-caps section headers with a rule beneath;
navy sub-headers; a navy banner at the top and foot of every page. Sizes are the
MAGNET spec's values read as Word half-points (48 = 24 pt), matching the sample report.
"""

from __future__ import annotations

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

from s2_watch.s2_report import CLOSING, HEADER_FIELDS, HEADER_LABELS, TOP_BANNER, S2Report

NAVY = "0D1B2A"
BLUE = "1565C0"
OFF_WHITE = "F5F7FA"
BORDER = "CCCCCC"
LEVEL_FILL = {"HIGH": ("FF0000", "000000"), "MODERATE": ("FFFF00", "000000"), "LOW": ("1B5E20", "FFFFFF")}
TABLE_WIDTH = 10080
# MagCon cell colour by the level's word (the scale's numbering isn't assumed).
MAGCON_FILL = {"SEVERE": ("C62828", "FFFFFF"), "CRITICAL": ("C62828", "FFFFFF"), "HIGH": ("EF6C00", "FFFFFF"),
               "ELEVATED": ("FFEB3B", "000000"), "GUARDED": ("1565C0", "FFFFFF"), "LOW": ("2E7D32", "FFFFFF"),
               "NORMAL": ("2E7D32", "FFFFFF")}


def _half(points: int) -> Pt:
    return Pt(points / 2)


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")          # ShadingType.CLEAR, never SOLID
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _cell_width(cell, dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = OxmlElement("w:tcW")
    width.set(qn("w:w"), str(dxa))
    width.set(qn("w:type"), "dxa")
    properties.append(width)


def _borders(table, color: str | None) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        if color:
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), "4")
            edge.set(qn("w:color"), color)
        else:
            edge.set(qn("w:val"), "nil")
        borders.append(edge)
    properties.append(borders)


def _table(container, rows: int, widths: list[int], border: str | None = BORDER):
    try:
        table = container.add_table(rows=rows, cols=len(widths))
    except TypeError:   # page headers and footers need the width up front
        table = container.add_table(rows, len(widths), Twips(TABLE_WIDTH))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    properties = table._tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(TABLE_WIDTH))
    width.set(qn("w:type"), "dxa")
    properties.append(width)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    properties.append(layout)
    margins = OxmlElement("w:tblCellMar")   # the sample's padding inside each cell
    for side, dxa in (("top", 70), ("bottom", 70), ("left", 110), ("right", 110)):
        margin = OxmlElement(f"w:{side}")
        margin.set(qn("w:w"), str(dxa))
        margin.set(qn("w:type"), "dxa")
        margins.append(margin)
    properties.append(margins)
    grid = table._tbl.tblGrid
    for column, dxa in zip(grid.findall(qn("w:gridCol")), widths):
        column.set(qn("w:w"), str(dxa))
    for row in table.rows:
        for cell, dxa in zip(row.cells, widths):
            _cell_width(cell, dxa)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _borders(table, border)
    return table


def _text(paragraph, words: str, size: int = 17, bold: bool = False, color: str = "000000",
          italic: bool = False) -> None:
    run = paragraph.add_run(words)
    run.font.name = "Arial"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    run.font.size = _half(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def _link(paragraph, url: str, size: int = 15) -> None:
    part = paragraph.part
    relation = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                              is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    for tag, value in (("w:rFonts", None), ("w:color", BLUE), ("w:sz", str(size)), ("w:u", "single")):
        element = OxmlElement(tag)
        if tag == "w:rFonts":
            element.set(qn("w:ascii"), "Arial")
            element.set(qn("w:hAnsi"), "Arial")
        else:
            element.set(qn("w:val"), value)
        properties.append(element)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = url
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _spacing(paragraph, before: int = 0, after: int = 4) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)


def _banner(container, words: str, size: int) -> None:
    table = _table(container, 1, [TABLE_WIDTH], border=None)
    cell = table.rows[0].cells[0]
    _shade(cell, NAVY)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(paragraph, 1, 1)
    _text(paragraph, words, size, bold=size >= 16, color="FFFFFF")


def _section_header(document, name: str) -> None:
    paragraph = document.add_paragraph()
    _spacing(paragraph, 10, 4)
    _text(paragraph, name.upper(), 26, bold=True, color=BLUE)
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", BLUE)):
        bottom.set(qn(key), value)
    borders.append(bottom)
    properties.append(borders)


def write_docx(report: S2Report, path: str) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Twips(12240), Twips(15840)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Twips(1080))
    style = document.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = _half(17)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Arial")
    style.paragraph_format.space_after = Pt(0)       # compact table rows; body paragraphs set their own
    style.paragraph_format.line_spacing = 1.0
    section.header_distance = Twips(400)             # room between the page banners and the text
    section.footer_distance = Twips(400)

    header = section.header
    header.paragraphs[0].text = ""
    _banner(header, TOP_BANNER, 16)
    footer = section.footer
    footer.paragraphs[0].text = ""
    _banner(footer, f"CLASSIFICATION: UNCLASSIFIED // OSINT   |   PREPARED BY: MAGNET S2 OSINT TEAM   |   "
                    f"DTG: {report.dtg}   |   www.magnethf.com", 15)

    # Title block: one navy cell.
    title = _table(document, 1, [TABLE_WIDTH], border=None).rows[0].cells[0]
    _shade(title, NAVY)
    lines = (("MAGNET S2", 48, True, "FFFFFF"), ("INTELLIGENCE REPORT", 24, False, "90CAF9"),
             (f"DTG: {report.dtg}  |  Geographic Focus: {report.fields.get('GEOGRAPHIC FOCUS', '')}", 18, False, "B0BEC5"),
             ("www.magnethf.com", 18, False, "64B5F6"))
    for index, (words, size, bold, color) in enumerate(lines):
        paragraph = title.paragraphs[0] if index == 0 else title.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _spacing(paragraph, 4 if index == 0 else 0, 4 if index == len(lines) - 1 else 0)
        _text(paragraph, words, size, bold, color)

    _section_header(document, "Report Identification")
    rows = [key for key in HEADER_FIELDS if key in report.fields]
    table = _table(document, len(rows), [2200, 7880])
    for row, key in zip(table.rows, rows):
        label, value = row.cells
        _shade(label, NAVY)
        _text(label.paragraphs[0], HEADER_LABELS[key], 17, bold=True, color="FFFFFF")
        words = report.fields[key]
        fill, ink = OFF_WHITE, "000000"
        if key == "MAGCON STATUS":
            level = next((word for word in MAGCON_FILL if word in words.upper()), None)
            if level:
                fill, ink = MAGCON_FILL[level]
        _shade(value, fill)
        _text(value.paragraphs[0], words, 17, bold=key == "MAGCON STATUS", color=ink)

    for name, blocks in report.sections:
        _section_header(document, name)
        for block in blocks:
            _block(document, block)

    closing = document.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(closing, 12, 0)
    _text(closing, CLOSING, 16, color="607D8B")
    document.save(path)   # the WordPress tags stay with the draft, not in the printed report


def _block(document, block) -> None:
    if block.kind == "subheader":
        paragraph = document.add_paragraph()
        _spacing(paragraph, 6, 2)
        _text(paragraph, block.text, 20, bold=True, color=NAVY)
    elif block.kind in ("paragraph", "note", "prompt"):
        paragraph = document.add_paragraph()
        _spacing(paragraph, 0, 4)
        quiet = block.kind != "paragraph"
        _text(paragraph, block.text, 16 if quiet else 17, italic=quiet, color="607D8B" if quiet else "000000")
    elif block.kind == "bullets":
        for item in block.rows:
            paragraph = document.add_paragraph(style="List Bullet")
            _spacing(paragraph, 0, 2)
            _text(paragraph, item)
    elif block.kind == "timeline":
        table = _table(document, len(block.rows), [2000, 8080])
        for row, (when, event) in zip(table.rows, block.rows):
            _shade(row.cells[0], NAVY)
            row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _text(row.cells[0].paragraphs[0], when, 16, bold=True, color="FFFFFF")
            _shade(row.cells[1], OFF_WHITE)
            _text(row.cells[1].paragraphs[0], event, 16)
        _spacing(document.add_paragraph(), 0, 2)
    elif block.kind == "probability":
        table = _table(document, len(block.rows) + 1, [7000, 3080])
        for cell, words in zip(table.rows[0].cells, ("ASSESSMENT CATEGORY", "PROBABILITY")):
            _shade(cell, NAVY)
            _text(cell.paragraphs[0], words, 16, bold=True, color="FFFFFF")
        table.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for row, (category, level) in zip(table.rows[1:], block.rows):
            _shade(row.cells[0], OFF_WHITE)
            _text(row.cells[0].paragraphs[0], category, 16, bold=True)
            fill, ink = LEVEL_FILL.get(level.upper(), (OFF_WHITE, "000000"))
            _shade(row.cells[1], fill)
            row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _text(row.cells[1].paragraphs[0], level.upper(), 16, bold=True, color=ink)
        _spacing(document.add_paragraph(), 0, 2)
    elif block.kind == "sources":
        for number, name, rating, description, url in block.rows:
            paragraph = document.add_paragraph()
            _spacing(paragraph, 0, 2)
            _text(paragraph, f"[{number}]  ", 15, bold=True)
            _text(paragraph, name, 15, bold=True)
            if rating:
                _text(paragraph, f" {rating}", 15, bold=True, color=NAVY)
            if description:
                _text(paragraph, f" — {description}", 15, color="607D8B")
            if url:
                _text(paragraph, " — ", 15, color="607D8B")
                _link(paragraph, url)
