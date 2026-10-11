from datetime import datetime, timezone

from docx import Document

from s2_watch.analysis.base import AnalysisResult
from s2_watch.followup import FollowUpFacts
from s2_watch.locating import ResolvedLocation
from s2_watch.models import Message
from s2_watch.s2_docx import write_docx
from s2_watch.s2_html import render_html
from s2_watch.s2_report import (
    SECTIONS, admiralty, file_stem, fill_assessment, followup_report, js8_report, magnet_dtg, parse_report,
    period_label,
)

NOW = datetime(2026, 10, 10, 15, 30, tzinfo=timezone.utc)
MEMBERS = [
    {"first_seen": "2026-10-09T23:05:00+00:00", "source": "MagCommstat", "sender": "KF5ABC", "summary": "Flooding — Tulsa, OK",
     "text": "TULSA OK - RIVER FLOODING, 2 ROADS CLOSED", "corroboration": "", "urgency": "high", "fields": {}},
    {"first_seen": "2026-10-10T01:40:00+00:00", "source": "OSINT", "sender": "FEMA", "summary": "Flood declaration — Oklahoma",
     "text": "FEMA: flood declaration Oklahoma", "corroboration": "", "urgency": "medium",
     "fields": {"url": "https://www.fema.gov/disaster/4999"}},
]
FACTS = FollowUpFacts(
    Message("MagCommstat", MEMBERS[0]["text"], received_at=datetime(2026, 10, 9, 23, 5, tzinfo=timezone.utc),
            fields={"sender": "KF5ABC"}),
    AnalysisResult(True, "Climate / Environment", "high", "Flooding — Tulsa, OK", 0.9),
    [ResolvedLocation("city", 36.154, -95.993, "Tulsa, OK", "OK", 6)], region=6, state="OK",
    official="✓ NWS Flood Warning · radio 45 min earlier",
)


def test_magnet_dates_and_names() -> None:
    assert magnet_dtg(NOW) == "261010-1530Z"
    assert period_label(datetime(2026, 5, 29, tzinfo=timezone.utc), datetime(2026, 5, 30, tzinfo=timezone.utc)) == "29–30 May 2026"
    assert file_stem("Iran peace deal talks resume", "260614-1800Z") == "MAGNET_S2_IRAN_PEACE_DEAL_TALKS_260614-1800Z"
    assert admiralty("OSINT", "USGS", True, 2) == "[A1]" and admiralty("JS8Call", "W1AW", False, 1) == "[F6]"


def test_followup_report_has_the_magnet_structure_and_parses_back() -> None:
    text = followup_report(FACTS, MEMBERS, now=NOW, magcon="3 – ELEVATED (NO CHANGE)")
    report = parse_report(text)
    assert [name for name, _blocks in report.sections] == list(SECTIONS)
    assert report.fields["DTG"] == "261010-1530Z" and report.fields["PRECEDENCE"] == "IM – IMMEDIATE"
    assert report.fields["REPORTING PERIOD"] == "09–10 Oct 2026"
    assert report.fields["GEOGRAPHIC FOCUS"] == "Tulsa, OK, MARS Region 6"
    situation = dict(report.sections)["SITUATION"]
    timeline = next(block for block in situation if block.kind == "timeline")
    assert timeline.rows[0] == ("~2305Z 09 Oct", "MagCommstat (KF5ABC): Flooding — Tulsa, OK")
    sources = next(block for block in dict(report.sections)["SOURCE LIST"] if block.kind == "sources")
    assert sources.rows[0][1:3] == ("MagCommstat (KF5ABC)", "[F1]")
    assert sources.rows[1][1:3] == ("FEMA", "[A1]") and sources.rows[1][4] == "https://www.fema.gov/disaster/4999"
    assert "corroborates" in text and 0 < len(report.tags) <= 15
    js8 = js8_report(text, "n1rog")
    assert js8.startswith("MAGNET S2 261010-1530Z IM FLOODING") and js8.endswith("DE N1ROG") and len(js8) <= 200


def test_word_and_web_exports_follow_the_format(tmp_path) -> None:
    report = parse_report(followup_report(FACTS, MEMBERS, now=NOW))
    path = tmp_path / f"{file_stem(report.subject, report.dtg)}.docx"
    write_docx(report, str(path))
    document = Document(str(path))
    words = "\n".join(p.text for p in document.paragraphs)
    cells = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    header = "\n".join(cell.text for table in document.sections[0].header.tables for row in table.rows for cell in row.cells)
    assert "SUMMARY (BLUF)" in words and "SOURCE LIST" in words and "www.MAGNETHF.COM" in words
    assert "MagCon Status" in cells and "~2305Z 09 Oct" in cells and "INTELLIGENCE REPORT" in cells
    assert header == "UNCLASSIFIED  //  OSINT  –  MAGNET S2 OSINT TEAM"
    assert "FOR OFFICIAL USE" not in words + cells + header
    section = document.sections[0]
    assert (section.page_width.twips, section.left_margin.twips) == (12240, 1080)
    html = render_html(report)
    assert "Summary (BLUF)" in html and "lvl-high" in html and "FOR OFFICIAL USE" not in html
    assert "https://www.fema.gov/disaster/4999" in html


def test_model_assessment_fills_the_prompt_and_keeps_the_table() -> None:
    text = followup_report(FACTS, MEMBERS, now=NOW)
    filled = fill_assessment(text, ["Reporting indicates sustained river flooding."], ["Are evacuations ordered?"])
    comments = dict(parse_report(filled).sections)["COMMENTS / ASSESSMENT"]
    assert comments[0].text == "Reporting indicates sustained river flooding."
    assert any(block.kind == "probability" for block in comments)
    gaps = dict(parse_report(filled).sections)["INTELLIGENCE GAPS"][0].rows
    assert gaps[-1] == "Are evacuations ordered?" and len(gaps) >= 2
